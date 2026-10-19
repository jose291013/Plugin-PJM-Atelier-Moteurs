#!/usr/bin/env python3
"""Prepare a Pressero product import file from a fresh export.

Usage:
  python3 prepare_import.py EXPORT.xlsx CHANGES.json OUTPUT.xlsx

CHANGES.json = list of row specs:
  [
    {"product_id": "187cd9dc-...", "set": {"Short Description": "..."}},
    {"copy_from": "187cd9dc-...", "create": true,
     "set": {"Product Name": "...", "Url Name": "...", "MIS Product ID price 1": "..."}}
  ]
- "product_id": update that existing row (only the given columns change).
- "copy_from" + "create": start from that row, empty Product Id and dates, apply "set".
Only the listed rows are kept. Image URLs lose their "?stamp=" query string.
Prints warnings for rows with 5 Site Group columns filled (possible truncation).
"""
import json
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")
DATE_COLS = ["Last modified date", "Creation date", "Modification date", "Deletion date"]


def main(src, changes_path, out):
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    rows = {}
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(r, col["Product Id"]).value
        if pid:
            rows[pid] = [ws.cell(r, c).value for c in range(1, len(hdr) + 1)]
    specs = json.load(open(changes_path, encoding="utf-8"))
    urls = {v[col["Url Name"] - 1] for v in rows.values()}
    out_rows = []
    for s in specs:
        base_id = s.get("product_id") or s.get("copy_from")
        if base_id not in rows:
            sys.exit(f"Product Id introuvable dans l'export : {base_id}")
        vals = list(rows[base_id])
        for k, v in s.get("set", {}).items():
            if k not in col:
                sys.exit(f"Colonne inconnue : {k}")
            vals[col[k] - 1] = v
        if s.get("create"):
            vals[col["Product Id"] - 1] = None
            for d in DATE_COLS:
                vals[col[d] - 1] = None
            url = vals[col["Url Name"] - 1]
            if url in urls:
                sys.exit(f"Url Name déjà utilisée : {url}")
        img = vals[col["Primary Image URL"] - 1]
        if isinstance(img, str) and "?" in img:
            vals[col["Primary Image URL"] - 1] = img.split("?")[0]
        groups = [vals[col[f"Site Group {i}"] - 1] for i in range(1, 6)]
        if all(groups):
            print(f"ATTENTION {vals[col['Product Name'] - 1]} : 5 groupes renseignés, "
                  "vérifier dans l'admin qu'il n'en a pas plus (troncature).")
        out_rows.append(vals)
    ws.delete_rows(2, ws.max_row)
    for i, vals in enumerate(out_rows, start=2):
        for c, v in enumerate(vals, start=1):
            ws.cell(i, c).value = v
    wb.save(out)
    print(f"OK : {len(out_rows)} ligne(s) -> {out}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:])
